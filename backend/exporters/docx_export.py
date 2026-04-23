"""
학습자료(material) 컴포넌트를 챕터별 .docx 파일로 변환.

각 챕터 docx는 6개 섹션(도입·핵심개념·GBB·실전시나리오·체크리스트·요약)을
heading + body 텍스트로 담아, 기존 aiworld-main의 Google Docs 대체물로 사용.
"""
from __future__ import annotations
from io import BytesIO
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH


KIND_BADGE = {
    "empathy_opener": "공감 오프너",
    "concept":        "핵심 개념",
    "example":        "Good / Better / Best",
    "template":       "실전 시나리오",
    "reflection":     "정리 / 체크리스트",
}


def build_material_docx(course_name: str, chapter_id: str, chapter_name: str, material: dict) -> bytes:
    """한 챕터의 학습자료를 docx 바이트로 반환."""
    doc = Document()

    # 전체 기본 스타일
    style = doc.styles["Normal"]
    style.font.name = "맑은 고딕"
    style.font.size = Pt(11)

    section = doc.sections[0]
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    # 헤더: 코스명(작게) + 챕터 타이틀(크게)
    p = doc.add_paragraph()
    r = p.add_run(course_name or "")
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x71, 0x71, 0x7A)

    title = doc.add_heading(level=1)
    t = title.add_run(f"[{chapter_id}] {chapter_name}")
    t.font.size = Pt(20)
    t.font.color.rgb = RGBColor(0x18, 0x18, 0x1B)

    # 섹션 반복
    sections = material.get("sections") or []
    for s in sections:
        heading = s.get("heading", "").strip()
        kind = s.get("kind", "")
        body = s.get("body", "") or ""

        # 섹션 제목
        h = doc.add_heading(level=2)
        run = h.add_run(heading)
        run.font.size = Pt(14)
        run.font.color.rgb = RGBColor(0x27, 0x27, 0x2A)

        # kind 배지
        if kind and kind in KIND_BADGE:
            sub = doc.add_paragraph()
            sub_run = sub.add_run(f"[{KIND_BADGE[kind]}]")
            sub_run.font.size = Pt(9)
            sub_run.font.color.rgb = RGBColor(0xA1, 0xA1, 0xAA)

        # 본문 — 줄바꿈 보존해서 문단 단위로 나눠 추가
        for para in body.split("\n"):
            if para.strip() == "":
                doc.add_paragraph()
                continue
            p = doc.add_paragraph()
            r = p.add_run(para)
            r.font.size = Pt(11)
            # 줄간격
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.space_after = Pt(4)

    # 푸터
    footer = doc.sections[0].footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = fp.add_run("AI World Harness — 자동 생성 학습자료")
    fr.font.size = Pt(8)
    fr.font.color.rgb = RGBColor(0xA1, 0xA1, 0xAA)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
